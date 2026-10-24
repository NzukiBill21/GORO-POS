"""Official shop onboarding template (3 example rows). Run from repo root:
  python tests/fixtures/shop-import/generate_shop_template.py
Writes docs/templates/ and public/templates/ for download in the app."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
OUT_DIRS = [
    os.path.join(ROOT, "docs", "templates"),
    os.path.join(ROOT, "public", "templates"),
]
FNAME = "GOROVIAN_Shop_Product_and_Stock_Template.xlsx"

PCOLS = [
    "Product name",
    "Product SKU",
    "Variant SKU",
    "Variant name",
    "Category",
    "Brand",
    "Unit",
    "Size",
    "Color",
    "Barcode",
    "Cost price",
    "Selling price",
    "Min stock",
    "Track stock",
]
PRODUCT_ROWS = [
    [
        "Oxford Cotton Shirt",
        "SR-0001",
        "SR-0001-WHT-M",
        "M / White",
        "Shirts",
        "Marvin Tailors",
        "pcs",
        "M",
        "White",
        "",
        900,
        1440,
        5,
        "Yes",
    ],
    [
        "Oxford Cotton Shirt",
        "SR-0001",
        "SR-0001-BLU-L",
        "L / Blue",
        "Shirts",
        "Marvin Tailors",
        "pcs",
        "L",
        "Blue",
        "",
        900,
        1440,
        5,
        "Yes",
    ],
    [
        "Everyday Court Sneakers",
        "SN-0001",
        "SN-0001-WHT-41",
        "41 / White",
        "Sneakers",
        "StrideX",
        "pair",
        "41",
        "White",
        "",
        2400,
        3600,
        3,
        "Yes",
    ],
]
STOCK_ROWS = [
    ["SR-0001-WHT-M", 12],
    ["SR-0001-BLU-L", 8],
    ["SN-0001-WHT-41", 6],
]

HF = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HFILL = PatternFill("solid", fgColor="1F3864")
BF = Font(name="Arial", size=10)
thin = Side(style="thin", color="BFBFBF")
BR = Border(left=thin, right=thin, top=thin, bottom=thin)


def table(ws, cols, rows, widths, text_cols=(), num_cols=(), int_cols=()):
    ws.append(cols)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = HF
        c.fill = HFILL
        c.border = BR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font = BF
            c.border = BR
            if c.column in text_cols:
                c.number_format = "@"
            if c.column in num_cols:
                c.number_format = "#,##0.00"
            if c.column in int_cols:
                c.number_format = "0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"


def text_sheet(ws, rows, w=(26, 110)):
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = w[0]
    ws.column_dimensions["B"].width = w[1]
    for row in ws.iter_rows():
        for c in row:
            c.font = BF
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(name="Arial", bold=True, size=12)


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "START HERE"
    text_sheet(
        ws,
        [
            ["GOROVIAN POS — Shop product & opening stock template"],
            [
                "Purpose",
                "Fill in PRODUCTS, then import opening quantities on OPENING STOCK. "
                "Use Products > Import Excel in Gorovian POS (Step 1 catalogue, Step 2 stock).",
            ],
            [
                "Order",
                "1) Import PRODUCTS (create catalogue).  2) Import OPENING STOCK (adds to default warehouse).",
            ],
            [
                "SKUs",
                "Product SKU groups variants under one product. Variant SKU must be unique per sellable line.",
            ],
            [
                "Barcodes",
                "Optional. Leave blank if you assign barcodes later in the app.",
            ],
            [
                "Sheets",
                "Only PRODUCTS and OPENING STOCK are read by the importer; this START HERE sheet is ignored.",
            ],
        ],
    )
    pw = [34, 14, 20, 24, 18, 18, 8, 12, 12, 16, 12, 13, 11, 11]
    pt = (2, 3, 8, 9, 10)
    w1 = wb.create_sheet("PRODUCTS")
    table(w1, PCOLS, PRODUCT_ROWS, pw, pt, (11, 12), (13,))
    w2 = wb.create_sheet("OPENING STOCK")
    table(w2, ["Variant SKU", "Quantity"], STOCK_ROWS, [24, 12], (1,), (), (2,))
    wb.active = 1
    return wb


def main():
    wb = build()
    for d in OUT_DIRS:
        os.makedirs(d, exist_ok=True)
        path = os.path.join(d, FNAME)
        wb.save(path)
        print("Wrote", path)


if __name__ == "__main__":
    main()
